"""
ingest.py
---------
Turns an uploaded file into a DuckDB database file that connectors.py can
query. Replaces the old pandas -> SQLite path.

WHY DUCKDB
  * Reads CSV / Parquet / JSON natively, in parallel, streaming straight from
    disk -- the file is never loaded into Python memory.
  * Columnar + vectorised: aggregations over 50M rows take about a second,
    where SQLite full-scans row by row.
  * Spills to disk when a query needs more RAM than `DUCKDB_MEMORY_LIMIT`.
  * One .duckdb file per session, so users can no longer overwrite each other.

Public API:
    load_file_to_dataset_config(file_path, db_path, table_name, progress=None) -> dict
returns a DATASET_CONFIG-shaped dict: {"type": "duckdb", "db_path": ..., ...}
"""

import json
import os
import re
from pathlib import Path
from typing import Callable, Optional

import duckdb

from config import DUCKDB_MEMORY_LIMIT, DUCKDB_THREADS, PROFILE_MAX_COLUMNS

TABLE_NAME = "uploaded_data"
Progress = Optional[Callable[[str], None]]   # callback(stage_text)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def _q(identifier: str) -> str:
    """Quote an SQL identifier."""
    return '"' + identifier.replace('"', '""') + '"'


def _lit(value: str) -> str:
    """Quote an SQL string literal (file paths)."""
    return "'" + value.replace("'", "''") + "'"


def _clean_name(name: str, taken: set) -> str:
    """'Order Date (UTC)' -> 'order_date_utc'. Keeps names unique and non-empty
    so the LLM never has to deal with quoting rules."""
    base = re.sub(r"[^0-9a-zA-Z]+", "_", str(name)).strip("_").lower() or "col"
    if base[0].isdigit():
        base = "c_" + base
    candidate, i = base, 2
    while candidate in taken:
        candidate = f"{base}_{i}"
        i += 1
    taken.add(candidate)
    return candidate


def _normalise_columns(con: duckdb.DuckDBPyConnection, table: str) -> None:
    cols = [r[0] for r in con.execute(f"DESCRIBE {_q(table)}").fetchall()]
    taken: set = set()
    renames = [(c, _clean_name(c, taken)) for c in cols]
    for old, new in renames:
        if old != new:
            # two-step rename avoids clashes like "A" -> "a" when "a" exists
            con.execute(f"ALTER TABLE {_q(table)} RENAME COLUMN {_q(old)} TO {_q('__tmp_' + new)}")
    for old, new in renames:
        if old != new:
            con.execute(f"ALTER TABLE {_q(table)} RENAME COLUMN {_q('__tmp_' + new)} TO {_q(new)}")


# ---------------------------------------------------------------------------
# per-format loaders (all write into `table` inside `con`)
# ---------------------------------------------------------------------------
def _load_csv(con, path: str, table: str, progress: Progress, delimiter: Optional[str]) -> None:
    delim = f", delim={_lit(delimiter)}" if delimiter else ""
    # Try the fast path first; on a type-detection failure escalate:
    #   1. sample the whole file for types   2. read everything as text.
    attempts = [
        ("Reading file", f"read_csv({_lit(path)}, header=true, sample_size=100000, ignore_errors=false{delim})"),
        ("Re-scanning file to detect column types", f"read_csv({_lit(path)}, header=true, sample_size=-1{delim})"),
        ("Loading as text (mixed-type columns)", f"read_csv({_lit(path)}, header=true, all_varchar=true{delim})"),
    ]
    last_err = None
    for stage, source in attempts:
        if progress:
            progress(stage)
        try:
            con.execute(f"CREATE OR REPLACE TABLE {_q(table)} AS SELECT * FROM {source}")
            return
        except duckdb.Error as exc:
            last_err = exc
    raise ValueError(f"Could not parse the CSV: {last_err}")


def _load_parquet(con, path: str, table: str, progress: Progress) -> None:
    if progress:
        progress("Reading Parquet")
    con.execute(f"CREATE OR REPLACE TABLE {_q(table)} AS SELECT * FROM read_parquet({_lit(path)})")


def _load_json(con, path: str, table: str, progress: Progress) -> None:
    if progress:
        progress("Reading JSON")
    con.execute(f"CREATE OR REPLACE TABLE {_q(table)} AS SELECT * FROM read_json_auto({_lit(path)}, maximum_object_size=104857600)")


def _load_excel(con, path: str, table: str, progress: Progress) -> None:
    """Excel is capped at ~1M rows by the format itself. .xlsx is streamed
    row-by-row (openpyxl read-only mode) in batches so memory stays flat."""
    import pandas as pd

    if path.lower().endswith(".xls"):
        if progress:
            progress("Reading Excel (.xls)")
        df = pd.read_excel(path)
        con.register("_xl", df)
        con.execute(f"CREATE OR REPLACE TABLE {_q(table)} AS SELECT * FROM _xl")
        con.unregister("_xl")
        return

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise ValueError("The Excel sheet is empty.")
    names = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(header)]

    batch, total, first = [], 0, True
    BATCH = 50_000

    def flush():
        nonlocal batch, first
        if not batch:
            return
        df = pd.DataFrame(batch, columns=names)
        con.register("_xl", df)
        if first:
            con.execute(f"CREATE OR REPLACE TABLE {_q(table)} AS SELECT * FROM _xl")
            first = False
        else:
            con.execute(f"INSERT INTO {_q(table)} SELECT * FROM _xl")
        con.unregister("_xl")
        batch = []

    for row in rows:
        batch.append(row)
        total += 1
        if len(batch) >= BATCH:
            flush()
            if progress:
                progress(f"Reading Excel: {total:,} rows")
    flush()
    wb.close()
    if first:
        raise ValueError("The Excel sheet has a header but no data rows.")


# ---------------------------------------------------------------------------
# schema profile -- what the LLM actually reads
# ---------------------------------------------------------------------------
_NUMERIC = ("INT", "DECIMAL", "DOUBLE", "FLOAT", "REAL", "HUGEINT", "UBIGINT")
_TEMPORAL = ("DATE", "TIMESTAMP", "TIME")


def _build_profile(con, table: str, total_rows: int, progress: Progress) -> str:
    """One scan collects min / max / approximate distinct count per column.
    Low-cardinality text columns also get their value list. This lets the
    agent write correct filters and know when a question is ambiguous
    (e.g. 'sales in a region' when 5 regions exist) WITHOUT scanning the
    table itself while chatting."""
    if progress:
        progress("Profiling columns")
    cols = con.execute(f"DESCRIBE {_q(table)}").fetchall()   # name, type, ...
    profiled = cols[:PROFILE_MAX_COLUMNS]

    exprs = []
    for name, ctype, *_ in profiled:
        exprs.append(f"approx_count_distinct({_q(name)})")
        if any(t in ctype.upper() for t in _NUMERIC + _TEMPORAL + ("VARCHAR",)) and "[" not in ctype:
            exprs.append(f"min({_q(name)})::VARCHAR")
            exprs.append(f"max({_q(name)})::VARCHAR")
        else:
            exprs.append("NULL")
            exprs.append("NULL")
    stats = con.execute(f"SELECT {', '.join(exprs)} FROM {_q(table)}").fetchone()

    lines = [f"Table: {table}  ({total_rows:,} rows, {len(cols)} columns)"]
    for i, (name, ctype, *_) in enumerate(profiled):
        distinct, lo, hi = stats[3 * i], stats[3 * i + 1], stats[3 * i + 2]
        detail = f"~{distinct:,} distinct"
        up = ctype.upper()
        if "VARCHAR" in up and distinct is not None and distinct <= 25:
            vals = [r[0] for r in con.execute(
                f"SELECT DISTINCT {_q(name)} FROM {_q(table)} WHERE {_q(name)} IS NOT NULL ORDER BY 1 LIMIT 25").fetchall()]
            detail += "; values: " + ", ".join(repr(v) for v in vals)
        elif lo is not None and hi is not None and "VARCHAR" not in up:
            detail += f"; range: {lo} .. {hi}"
        lines.append(f"  - {name} ({ctype}) [{detail}]")
    if len(cols) > len(profiled):
        rest = ", ".join(f"{c[0]} ({c[1]})" for c in cols[len(profiled):])
        lines.append(f"  - (not profiled) {rest}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# main entry point
# ---------------------------------------------------------------------------
def load_file_to_dataset_config(
    file_path: str,
    db_path: str,
    table_name: str = TABLE_NAME,
    progress: Progress = None,
    original_name: Optional[str] = None,
) -> dict:
    """Loads `file_path` into a fresh DuckDB file at `db_path` and returns a
    DATASET_CONFIG-shaped dict. Raises ValueError for unsupported/corrupt
    files so the caller can show the user a clear message."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    shown_name = original_name or os.path.basename(file_path)
    ext = os.path.splitext(shown_name)[1].lower()
    Path(db_path).parent.mkdir(parents=True, exist_ok=True)
    if os.path.exists(db_path):
        os.remove(db_path)

    spill_dir = str(Path(db_path).parent / "spill")
    con = duckdb.connect(db_path, config={
        "memory_limit": DUCKDB_MEMORY_LIMIT,
        "threads": DUCKDB_THREADS,
        "temp_directory": spill_dir,
        "preserve_insertion_order": False,     # lets big loads run in parallel with less memory
    })
    try:
        if ext == ".csv":
            _load_csv(con, file_path, table_name, progress, None)
        elif ext == ".tsv":
            _load_csv(con, file_path, table_name, progress, "\t")
        elif ext in (".xlsx", ".xls"):
            _load_excel(con, file_path, table_name, progress)
        elif ext in (".json", ".jsonl", ".ndjson"):
            _load_json(con, file_path, table_name, progress)
        elif ext == ".parquet":
            _load_parquet(con, file_path, table_name, progress)
        else:
            raise ValueError(f"Unsupported file type '{ext}'. Supported: .csv, .tsv, .xlsx, .xls, .json, .jsonl, .parquet")

        total_rows = con.execute(f"SELECT count(*) FROM {_q(table_name)}").fetchone()[0]
        if total_rows == 0:
            raise ValueError(f"'{shown_name}' has no data rows.")

        if progress:
            progress("Cleaning column names")
        _normalise_columns(con, table_name)

        profile = _build_profile(con, table_name, total_rows, progress)
        columns = [r[0] for r in con.execute(f"DESCRIBE {_q(table_name)}").fetchall()]
        con.execute("CHECKPOINT")
    except duckdb.Error as exc:
        raise ValueError(f"Couldn't load '{shown_name}': {exc}") from exc
    finally:
        con.close()

    # The profile is stored next to the database so connectors.py can hand it
    # to the LLM instantly, with no scan, on every question.
    Path(db_path).with_suffix(".profile.txt").write_text(profile, encoding="utf-8")

    return {
        "type": "duckdb",
        "db_path": db_path,
        "target": table_name,
        "rows": total_rows,
        "description": (
            f"User-uploaded dataset from '{shown_name}' "
            f"({total_rows:,} rows, {len(columns)} columns: {', '.join(columns[:8])}"
            f"{', ...' if len(columns) > 8 else ''})."
        ),
    }
